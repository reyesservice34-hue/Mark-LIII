"""
Echte Dateien erzeugen — PDF, Excel, einfache Grafiken — statt nur Text, der so benannt ist.

Bewusste Grenze bei Bildern: Hier wird keine KI-Bildgenerierung angeboten, weil kein
Bildgenerierungs-Anbieter (DALL·E, Imagen, Stable Diffusion) in diesem System eingerichtet ist —
`image.create` zeichnet einfache, gestaltete Grafiken (Titel, Text, Formen, Farben), keine
fotorealistischen Bilder. Wer das will, muss zuerst einen Anbieter einrichten lassen; das hier
tut nicht so, als könnte es das schon.
"""
from __future__ import annotations

import io
import re


class RenderError(Exception):
    pass


# ── PDF ──────────────────────────────────────────────────────────────────────
def markdown_to_pdf(title: str, content: str) -> bytes:
    """Ein einfaches, sauber formatiertes PDF aus Titel + Text mit sehr einfacher Markdown-artiger
    Gliederung: `# `/`## ` werden Überschriften, `- `/`* ` werden Aufzählungspunkte, Leerzeilen
    trennen Absätze. Kein voller Markdown-Parser — genug für Angebote, Berichte, Briefe."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as e:
        raise RenderError("reportlab ist nicht installiert — PDF-Erzeugung ist nicht verfügbar.") from e

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], spaceAfter=10)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], spaceAfter=8)
    body = ParagraphStyle("Body", parent=styles["Normal"], spaceAfter=6, leading=15)

    def esc(t: str) -> str:
        return (t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=22 * mm, bottomMargin=20 * mm,
                            leftMargin=20 * mm, rightMargin=20 * mm, title=title)
    story = [Paragraph(esc(title), h1), Spacer(1, 4 * mm)]
    bullets: list = []

    def flush_bullets():
        nonlocal bullets
        if bullets:
            story.append(ListFlowable([ListItem(Paragraph(esc(b), body)) for b in bullets],
                                      bulletType="bullet", leftIndent=14))
            bullets = []

    for line in content.splitlines():
        s = line.rstrip()
        if not s.strip():
            flush_bullets()
            continue
        if s.startswith("## "):
            flush_bullets(); story.append(Paragraph(esc(s[3:].strip()), h2))
        elif s.startswith("# "):
            flush_bullets(); story.append(Paragraph(esc(s[2:].strip()), h1))
        elif re.match(r"^[-*]\s+", s):
            bullets.append(re.sub(r"^[-*]\s+", "", s))
        else:
            flush_bullets(); story.append(Paragraph(esc(s.strip()), body))
    flush_bullets()
    doc.build(story)
    return buf.getvalue()


# ── Excel ────────────────────────────────────────────────────────────────────
def rows_to_excel(sheets: list[dict]) -> bytes:
    """`sheets`: [{"name": "Tabelle1", "headers": [...], "rows": [[...], ...]}, ...] → echte .xlsx-Bytes."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise RenderError("openpyxl ist nicht installiert — Excel-Erzeugung ist nicht verfügbar.") from e
    if not sheets:
        raise RenderError("Mindestens ein Arbeitsblatt wird gebraucht.")

    wb = Workbook()
    wb.remove(wb.active)
    for i, sheet in enumerate(sheets):
        name = str(sheet.get("name") or f"Tabelle{i + 1}")[:31] or f"Tabelle{i + 1}"
        ws = wb.create_sheet(name)
        headers = [str(h) for h in (sheet.get("headers") or [])]
        if headers:
            ws.append(headers)
            for c in ws[1]:
                c.font = Font(bold=True)
        for row in sheet.get("rows") or []:
            ws.append([("" if v is None else v) for v in row])
        widths: dict[int, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                widths[cell.column] = max(widths.get(cell.column, 8), min(60, len(str(cell.value)) + 2))
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Einfache Grafiken ────────────────────────────────────────────────────────
def make_image(*, title: str = "", lines: list[str] | None = None, width: int = 1200, height: int = 800,
              bg: str = "#141619", fg: str = "#f7f2ec", accent: str = "#ffb56e") -> bytes:
    """Eine einfache, gestaltete Grafik (Titel + Zeilen auf Hintergrundfarbe) — kein Foto, keine KI-Kunst.

    Für echte Bild-/Fotoerzeugung müsste ein Bildgenerierungs-Anbieter eingerichtet werden; das
    gibt es hier nicht, also wird auch nicht so getan, als könnte dieses Werkzeug das.
    """
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError as e:
        raise RenderError("Pillow ist nicht installiert — Bilderzeugung ist nicht verfügbar.") from e

    width = max(200, min(width, 4000))
    height = max(200, min(height, 4000))
    img = Image.new("RGB", (width, height), bg)
    draw = ImageDraw.Draw(img)

    def load_font(size: int):
        for path in ("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
                    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
        return ImageFont.load_default()

    pad = int(width * 0.06)
    draw.rectangle([pad // 2, pad // 2, width - pad // 2, height - pad // 2], outline=accent, width=3)
    y = pad
    if title:
        f_title = load_font(max(24, width // 18))
        draw.text((pad, y), title, font=f_title, fill=fg)
        bbox = draw.textbbox((pad, y), title, font=f_title)
        y = bbox[3] + int(pad * 0.6)
        draw.line([(pad, y), (width - pad, y)], fill=accent, width=2)
        y += int(pad * 0.5)
    f_body = load_font(max(16, width // 32))
    for line in (lines or []):
        draw.text((pad, y), str(line), font=f_body, fill=fg)
        bbox = draw.textbbox((pad, y), str(line), font=f_body)
        y += (bbox[3] - bbox[1]) + int(pad * 0.3)
        if y > height - pad:
            break
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()
